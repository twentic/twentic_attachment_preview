# -*- coding: utf-8 -*-
import io
import json
import logging

from odoo import http
from odoo.http import request

_logger = logging.getLogger(__name__)

# Safety limits to prevent OOM with very large spreadsheets
_MAX_ROWS   = 2000
_MAX_SHEETS = 20


class ChatterImagePreviewController(http.Controller):

    @http.route(
        '/chatter_image_preview/xlsx/<int:attachment_id>',
        type='http',
        auth='user',
        methods=['GET'],
    )
    def xlsx_preview(self, attachment_id, **kwargs):
        """
        Read an XLSX attachment and return its sheets as JSON rows.

        Response format:
            { "sheets": [ { "name": "Sheet1", "rows": [["A1","B1"], ...] } ] }
        On error:
            { "error": "<message>" }

        Uses openpyxl which ships with every Odoo 18 installation.
        Limits: max 2000 rows and 20 sheets to keep responses lightweight.
        """
        attachment = request.env['ir.attachment'].browse(attachment_id)
        if not attachment.exists():
            return self._json({'error': 'Attachment not found.'})

        try:
            import openpyxl
        except ImportError:
            return self._json({
                'error': 'Python package "openpyxl" is not available on this server.'
            })

        file_bytes = attachment.raw
        if not file_bytes:
            return self._json({'error': 'The attachment has no content.'})

        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(file_bytes),
                read_only=True,
                data_only=True,
            )

            sheets = []
            for sheet_name in wb.sheetnames[:_MAX_SHEETS]:
                ws = wb[sheet_name]
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= _MAX_ROWS:
                        break
                    rows.append([
                        '' if cell is None else str(cell)
                        for cell in row
                    ])
                sheets.append({'name': sheet_name, 'rows': rows})

            wb.close()
            return self._json({'sheets': sheets})

        except Exception:
            _logger.exception(
                'chatter_image_preview: could not parse XLSX attachment %d',
                attachment_id,
            )
            return self._json({'error': 'Could not parse the file. It may be corrupt or password-protected.'})

    @http.route(
        '/chatter_image_preview/docx/<int:attachment_id>',
        type='http',
        auth='user',
        methods=['GET'],
    )
    def docx_preview(self, attachment_id, **kwargs):
        """
        Convert a DOCX attachment to HTML and return it as JSON.

        Response format:
            { "html": "<html string>" }
        On error:
            { "error": "<message>" }

        Uses mammoth for DOCX-to-HTML conversion (must be pip-installed).
        Only .docx (Open XML) format is supported; legacy .doc files will
        get an informative error message.
        """
        attachment = request.env['ir.attachment'].browse(attachment_id)
        if not attachment.exists():
            return self._json({'error': 'Attachment not found.'})

        try:
            import mammoth
        except ImportError:
            return self._json({
                'error': 'Python package "mammoth" is not available on this server. '
                         'Ask your administrator to run: pip install mammoth'
            })

        file_bytes = attachment.raw
        if not file_bytes:
            return self._json({'error': 'The attachment has no content.'})

        # Legacy .doc (binary) files are not supported by mammoth
        name = (attachment.name or attachment.datas_fname or '').lower()
        if name.endswith('.doc') and not name.endswith('.docx'):
            return self._json({
                'error': 'Legacy .doc files cannot be previewed. '
                         'Please convert the file to .docx format first.'
            })

        try:
            result = mammoth.convert_to_html(io.BytesIO(file_bytes))
            return self._json({'html': result.value})
        except Exception:
            _logger.exception(
                'chatter_image_preview: could not convert DOCX attachment %d',
                attachment_id,
            )
            return self._json({
                'error': 'Could not convert the file. It may be corrupt or password-protected.'
            })

    @staticmethod
    def _json(data):
        return request.make_response(
            json.dumps(data),
            headers=[('Content-Type', 'application/json')],
        )
